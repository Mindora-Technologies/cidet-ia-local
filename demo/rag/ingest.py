#!/usr/bin/env python3
"""Ingesta del corpus a Qdrant.

Pipeline complet: extracció → fragmentació → embeddings → índex vectorial.
Cada pas és configurable perquè a la C4 es puguin comparar les alternatives.

Ús:
    python ingest.py --reset                    # recrea la col·lecció i ingesta
    python ingest.py --extractor docling        # extracció de qualitat (lenta)
    python ingest.py --chunk-size 500 --overlap 50
    python ingest.py --dry-run                  # només mostra els fragments

Variables d'entorn (vegeu demo/.env.example):
    QDRANT_URL, QDRANT_COLLECTION, OLLAMA_URL, EMBED_MODEL, CORPUS_DIR
"""
from __future__ import annotations

import argparse
import hashlib
import os
import re
import sys
import time
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterator, Literal

import httpx

# ---------------------------------------------------------------- config
ARREL = Path(__file__).resolve().parents[2]
CORPUS_DIR = Path(os.getenv("CORPUS_DIR", ARREL / "demo" / "corpus"))
QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
COLLECTION = os.getenv("QDRANT_COLLECTION", "distribucions_valles")
OLLAMA_URL = os.getenv("OLLAMA_URL", "http://localhost:11434")
EMBED_MODEL = os.getenv("EMBED_MODEL", "qwen3-embedding:0.6b")

# ~800 tokens amb 100 de solapament. En català comptem ~3,2 car./token:
# tokenitza pitjor que l'anglès, i per això el mateix text ocupa més context.
CAR_PER_TOKEN = 3.2

TIPUS_DOC = {
    "manual-garanties": "política de garanties",
    "procediment-devolucions": "procediment",
    "tarifes": "tarifa de preus",
    "fitxa-producte": "fitxa tècnica",
    "politica-enviaments": "política d'enviaments",
    "condicions-generals-venda": "condicions contractuals",
    "condicions-pagament": "condicions contractuals",
    "manual-magatzem": "manual intern",
    "cataleg-serveis": "catàleg",
    "homologacio-proveidors": "procediment",
    "faq-servei-tecnic": "preguntes freqüents",
    "acta-qualitat": "acta de reunió",
}


@dataclass
class Fragment:
    """Un tros de document, llest per vectoritzar."""
    text: str
    fitxer: str
    pagina: int
    tipus_document: str
    data_document: str
    posicio: int
    id: str = field(default="")

    def __post_init__(self) -> None:
        if not self.id:
            llavor = f"{self.fitxer}:{self.pagina}:{self.posicio}"
            self.id = str(uuid.UUID(hashlib.sha256(llavor.encode()).hexdigest()[:32]))


# ------------------------------------------------------------- extracció
def extreu_pypdf(cami: Path) -> list[tuple[int, str]]:
    """Ràpid i suficient per a text corregut. No entén taules ni columnes."""
    from pypdf import PdfReader
    return [(i, (p.extract_text() or ""))
            for i, p in enumerate(PdfReader(cami).pages, start=1)]


def extreu_docling(cami: Path) -> list[tuple[int, str]]:
    """Millor qualitat: manté l'estructura i converteix les taules a Markdown.

    És molt més lent i baixa models la primera vegada. A la C4 es comparen
    els dos extractors sobre tarifes-2026.pdf, que va ple de taules.
    """
    try:
        from docling.document_converter import DocumentConverter
    except ImportError:
        print("  ! docling no està instal·lat (pip install docling); "
              "faig servir pypdf", file=sys.stderr)
        return extreu_pypdf(cami)
    doc = DocumentConverter().convert(str(cami)).document
    return [(1, doc.export_to_markdown())]


def extreu_docx(cami: Path) -> list[tuple[int, str]]:
    """Word: paràgrafs i taules. Word no té pàgines fins que es renderitza."""
    from docx import Document
    doc = Document(cami)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for taula in doc.tables:
        for fila in taula.rows:
            cel = [c.text.strip() for c in fila.cells]
            if any(cel):
                parts.append(" | ".join(cel))
    return [(1, "\n".join(parts))]


def extreu(cami: Path, extractor: Literal["pypdf", "docling"]) -> list[tuple[int, str]]:
    if cami.suffix.lower() == ".docx":
        return extreu_docx(cami)
    return extreu_docling(cami) if extractor == "docling" else extreu_pypdf(cami)


# ----------------------------------------------------------- fragmentació
def neteja(text: str) -> str:
    text = text.replace("\xad", "").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    # peus de pàgina repetits: no aporten res i embruten els embeddings
    text = re.sub(r"Distribucions Vallès SL · CIDET · IA en local[^\n]*", "", text)
    text = re.sub(r"\bpàg\. \d+\b", "", text)
    return text.strip()


def fragmenta(text: str, mida_tokens: int, solapament_tokens: int) -> Iterator[str]:
    """Talla per frases, no per caràcters: un fragment tallat a mitja frase
    dóna un embedding pitjor i una cita lletja."""
    mida = int(mida_tokens * CAR_PER_TOKEN)
    solap = int(solapament_tokens * CAR_PER_TOKEN)

    frases = re.split(r"(?<=[.!?:])\s+|\n{2,}", text)
    frases = [f.strip() for f in frases if f.strip()]

    actual: list[str] = []
    llarg = 0
    for f in frases:
        if llarg + len(f) > mida and actual:
            yield " ".join(actual)
            # cua de solapament: manté el context entre fragments veïns
            cua: list[str] = []
            n = 0
            for t in reversed(actual):
                if n + len(t) > solap:
                    break
                cua.insert(0, t)
                n += len(t)
            actual, llarg = cua, n
        actual.append(f)
        llarg += len(f) + 1
    if actual:
        yield " ".join(actual)


def tipus_de(nom: str) -> str:
    for clau, tipus in TIPUS_DOC.items():
        if nom.startswith(clau):
            return tipus
    return "document"


def data_de(nom: str) -> str:
    m = re.search(r"(20\d{2})", nom)
    return f"{m.group(1)}-01-01" if m else "2026-01-01"


# -------------------------------------------------------------- embeddings
class Embedder:
    """Embeddings via Ollama. Seqüencial a propòsit: a l'aula compartim GPU."""

    def __init__(self, url: str, model: str) -> None:
        self.url, self.model = url.rstrip("/"), model
        self.client = httpx.Client(timeout=120.0)
        self.dim: int | None = None

    def comprova(self) -> None:
        try:
            r = self.client.get(f"{self.url}/api/tags")
            r.raise_for_status()
        except Exception as e:
            raise SystemExit(
                f"✘ no puc parlar amb Ollama a {self.url}: {e}\n"
                f"  Comprova que el servei està engegat: systemctl status ollama"
            ) from e
        noms = {m["name"] for m in r.json().get("models", [])}
        if self.model not in noms and f"{self.model}:latest" not in noms:
            raise SystemExit(
                f"✘ el model d'embeddings «{self.model}» no hi és.\n"
                f"  Baixa'l amb:  ollama pull {self.model}\n"
                f"  Models disponibles: {', '.join(sorted(noms)) or '(cap)'}"
            )

    def vector(self, text: str) -> list[float]:
        r = self.client.post(f"{self.url}/api/embed",
                             json={"model": self.model, "input": text})
        r.raise_for_status()
        v = r.json()["embeddings"][0]
        if self.dim is None:
            self.dim = len(v)
        return v


# ----------------------------------------------------------------- Qdrant
class Qdrant:
    def __init__(self, url: str, colleccio: str) -> None:
        self.url, self.col = url.rstrip("/"), colleccio
        self.client = httpx.Client(timeout=60.0)

    def comprova(self) -> None:
        try:
            self.client.get(f"{self.url}/collections").raise_for_status()
        except Exception as e:
            raise SystemExit(
                f"✘ no puc parlar amb Qdrant a {self.url}: {e}\n"
                f"  Aixeca'l amb:  docker compose up -d qdrant"
            ) from e

    def existeix(self) -> bool:
        return self.client.get(f"{self.url}/collections/{self.col}").status_code == 200

    def esborra(self) -> None:
        self.client.delete(f"{self.url}/collections/{self.col}")

    def crea(self, dim: int) -> None:
        r = self.client.put(f"{self.url}/collections/{self.col}", json={
            "vectors": {"size": dim, "distance": "Cosine"},
            # HNSW: m = veïns per node, ef_construct = amplada de la cerca en indexar.
            # Més alt = índex millor i més lent de construir.
            "hnsw_config": {"m": 16, "ef_construct": 128},
            "optimizers_config": {"default_segment_number": 2},
        })
        r.raise_for_status()
        # Índexs de càrrega útil: permeten filtrar per tipus abans de la cerca vectorial.
        for camp, tipus in (("tipus_document", "keyword"), ("fitxer", "keyword")):
            self.client.put(f"{self.url}/collections/{self.col}/index",
                            json={"field_name": camp, "field_schema": tipus})

    def puja(self, punts: list[dict]) -> None:
        r = self.client.put(f"{self.url}/collections/{self.col}/points?wait=true",
                            json={"points": punts})
        r.raise_for_status()

    def compte(self) -> int:
        r = self.client.post(f"{self.url}/collections/{self.col}/points/count",
                             json={"exact": True})
        return r.json()["result"]["count"]


# ------------------------------------------------------------------- main
def barra(fet: int, total: int, etiqueta: str, ample: int = 34) -> None:
    plens = int(ample * fet / max(total, 1))
    pct = 100 * fet / max(total, 1)
    sys.stdout.write(f"\r  [{'█'*plens}{'·'*(ample-plens)}] {pct:5.1f}%  {etiqueta[:38]:38s}")
    sys.stdout.flush()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--corpus", type=Path, default=CORPUS_DIR)
    ap.add_argument("--extractor", choices=["pypdf", "docling"], default="pypdf",
                    help="pypdf és ràpid; docling entén taules però va lent")
    ap.add_argument("--chunk-size", type=int, default=800, help="tokens per fragment")
    ap.add_argument("--overlap", type=int, default=100, help="tokens de solapament")
    ap.add_argument("--reset", action="store_true", help="recrea la col·lecció des de zero")
    ap.add_argument("--dry-run", action="store_true", help="no toca Qdrant ni Ollama")
    ap.add_argument("--limit", type=int, default=0, help="processa només N documents")
    args = ap.parse_args()

    if not args.corpus.is_dir():
        raise SystemExit(f"✘ no trobo el corpus a {args.corpus}\n"
                         f"  Genera'l amb:  python demo/generar_corpus.py")

    fitxers = sorted(p for p in args.corpus.iterdir()
                     if p.suffix.lower() in (".pdf", ".docx"))
    if args.limit:
        fitxers = fitxers[:args.limit]
    if not fitxers:
        raise SystemExit(f"✘ no hi ha cap document a {args.corpus}")

    print(f"\n▶ Ingesta del corpus  ({len(fitxers)} documents)")
    print(f"  extractor={args.extractor}  fragments={args.chunk_size}t "
          f"solapament={args.overlap}t  model={EMBED_MODEL}\n")

    # ---------------------------------------------------- extreure i trossejar
    t0 = time.time()
    fragments: list[Fragment] = []
    sense_text: list[str] = []
    for i, f in enumerate(fitxers, start=1):
        barra(i - 1, len(fitxers), f.name)
        pagines = extreu(f, args.extractor)
        caracters = 0
        for num_pagina, brut in pagines:
            text = neteja(brut)
            caracters += len(text)
            for pos, tros in enumerate(fragmenta(text, args.chunk_size, args.overlap)):
                if len(tros) < 80:      # restes de peu de pàgina
                    continue
                fragments.append(Fragment(
                    text=tros, fitxer=f.name, pagina=num_pagina,
                    tipus_document=tipus_de(f.stem), data_document=data_de(f.stem),
                    posicio=pos,
                ))
        if caracters < 100:
            sense_text.append(f.name)
    barra(len(fitxers), len(fitxers), "fet")
    t_extraccio = time.time() - t0
    print(f"\n\n  {len(fragments)} fragments de {len(fitxers)} documents "
          f"en {t_extraccio:.1f} s")

    if sense_text:
        print(f"\n  ⚠ sense text extraïble: {', '.join(sense_text)}")
        print("    Són escanejos: sense OCR no entren al RAG. Ho veurem a la C4.")

    if fragments:
        mida_mitjana = sum(len(f.text) for f in fragments) / len(fragments)
        print(f"  mida mitjana: {mida_mitjana:.0f} caràcters "
              f"(~{mida_mitjana/CAR_PER_TOKEN:.0f} tokens)")

    if args.dry_run:
        print("\n--dry-run: mostra dels 3 primers fragments\n")
        for fr in fragments[:3]:
            print(f"  ── {fr.fitxer} p.{fr.pagina} #{fr.posicio} "
                  f"[{fr.tipus_document}]")
            print(f"     {fr.text[:220]}…\n")
        return

    # ------------------------------------------------------------- embeddings
    emb = Embedder(OLLAMA_URL, EMBED_MODEL)
    emb.comprova()
    qd = Qdrant(QDRANT_URL, COLLECTION)
    qd.comprova()

    print(f"\n▶ Embeddings amb {EMBED_MODEL}")
    t0 = time.time()
    vectors: list[list[float]] = []
    for i, fr in enumerate(fragments):
        barra(i, len(fragments), f"{fr.fitxer} p.{fr.pagina}")
        vectors.append(emb.vector(fr.text))
    barra(len(fragments), len(fragments), "fet")
    t_embed = time.time() - t0
    print(f"\n  dimensió del vector: {emb.dim}  ·  {t_embed:.1f} s "
          f"({len(fragments)/max(t_embed,0.01):.1f} fragments/s)")

    # ----------------------------------------------------------------- índex
    print(f"\n▶ Qdrant · col·lecció «{COLLECTION}»")
    if qd.existeix():
        if args.reset:
            qd.esborra()
            print("  col·lecció anterior esborrada (--reset)")
            qd.crea(emb.dim)
        else:
            print("  la col·lecció ja existeix; els punts amb el mateix id "
                  "se sobreescriuen (fes servir --reset per començar de zero)")
    else:
        qd.crea(emb.dim)
        print(f"  col·lecció creada (dim={emb.dim}, cosine, HNSW m=16)")

    t0 = time.time()
    LOT = 64
    for i in range(0, len(fragments), LOT):
        tall = fragments[i:i + LOT]
        qd.puja([
            {"id": fr.id, "vector": v,
             "payload": {k: v2 for k, v2 in asdict(fr).items() if k != "id"}}
            for fr, v in zip(tall, vectors[i:i + LOT])
        ])
        barra(min(i + LOT, len(fragments)), len(fragments), "pujant a Qdrant")
    print(f"\n  {qd.compte()} punts a la col·lecció  ·  {time.time()-t0:.1f} s")

    # --------------------------------------------------------------- resum
    per_tipus: dict[str, int] = {}
    for fr in fragments:
        per_tipus[fr.tipus_document] = per_tipus.get(fr.tipus_document, 0) + 1
    print("\n▶ Resum")
    print(f"  {'documents':<22} {len(fitxers)}")
    print(f"  {'fragments':<22} {len(fragments)}")
    print(f"  {'dimensió':<22} {emb.dim}")
    print(f"  {'temps total':<22} {t_extraccio + t_embed:.1f} s")
    for tipus, n in sorted(per_tipus.items(), key=lambda x: -x[1]):
        print(f"    · {tipus:<28} {n:>4}")
    print(f"\n✔ Llest. Prova-ho amb:  python demo/rag/agent.py "
          f"\"quina garantia tenen les eines elèctriques?\"\n")


if __name__ == "__main__":
    main()
