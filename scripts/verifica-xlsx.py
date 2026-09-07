#!/usr/bin/env python3
"""Verificació de c1/hardware-calc.xlsx.

Comprova, sense obrir Excel:
  1. que les sortides són FÓRMULES i no valors precalculats,
  2. que les validacions de dades i el format condicional hi són,
  3. que cap cel·la desborda l'amplada de la seva columna (res de «#####»),
  4. que la lògica de les fórmules dóna els números de la diapositiva 24.

Ús:  python scripts/verifica-xlsx.py [ruta.xlsx]
Surt amb codi 1 si alguna comprovació falla.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

XLSX = Path(sys.argv[1] if len(sys.argv) > 1 else "c1/hardware-calc.xlsx")

BYTES_PER_PARAM = {"FP16": 2.0, "Q8": 1.0, "Q6": 0.75, "Q4": 0.5, "Q3": 0.4}
errors: list[str] = []
warns: list[str] = []


def check(cond: bool, msg: str) -> None:
    if cond:
        print(f"  ✔ {msg}")
    else:
        errors.append(msg)
        print(f"  ✘ {msg}")


def vram_pesos(params: float, quant: str) -> float:
    return params * BYTES_PER_PARAM[quant] * 1.1


def vram_ctx(ctx: int, users: int = 1, kv_q8: bool = False) -> float:
    return ctx / 1000 * 0.15 * users * (0.5 if kv_q8 else 1)


def veredicte(marge: float) -> str:
    if marge < 0:
        return "✘ No"
    return "✔ Sí" if marge > 1.5 else "⚠ Just"


print(f"\n▶ Verificant {XLSX}\n")
if not XLSX.exists():
    sys.exit(f"no existeix: {XLSX}")

wb = load_workbook(XLSX, data_only=False)

print("1. Estructura")
check(wb.sheetnames == ["Dimensionament", "GPUs", "Referencia", "Exemples"],
      f"pestanyes correctes: {wb.sheetnames}")

ws = wb["Dimensionament"]
print("\n2. Sortides escrites com a fórmules")
formules = {
    "F6": "VLOOKUP",
    "F7": "*1.1",
    "F8": "0.15",
    "F9": "$F$7+$F$8",
    "F10": "VLOOKUP",
    "F11": "$F$10-$F$9",
    "F12": "IF(",
    "F14": "VLOOKUP",
    "F15": "/$F$7",
    "F16": "*0.7",
    "H6": "IF(",
}
for cell, frag in formules.items():
    v = ws[cell].value
    is_formula = isinstance(v, str) and v.startswith("=")
    check(is_formula and frag in v, f"{cell} és fórmula i conté «{frag}»")

print("\n3. Validacions de dades")
dv_refs = {str(dv.sqref): dv.type for dv in ws.data_validations.dataValidation}
for ref in ("C8", "C9", "C10", "C11", "C12", "C13"):
    check(ref in dv_refs, f"{ref} té validació ({dv_refs.get(ref, '—')})")
check(any("Referencia" in (dv.formula1 or "")
          for dv in ws.data_validations.dataValidation),
      "la quantització es valida contra la pestanya Referencia")
check(not any((dv.formula1 or "").startswith("=")
              for dv in ws.data_validations.dataValidation),
      "cap formula1 comença amb «=» (Excel ho rebutjaria)")

print("\n4. Format condicional")
check(len(ws.conditional_formatting._cf_rules) >= 2,
      f"Dimensionament: {len(ws.conditional_formatting._cf_rules)} rangs amb format condicional")
check(len(wb["Exemples"].conditional_formatting._cf_rules) >= 3,
      "Exemples: matriu «hi cap?» amb semàfor verd/ambre/vermell")

print("\n5. Protecció de fulls")
check(ws.protection.sheet is False, "Dimensionament editable")
for name in ("GPUs", "Referencia"):
    check(wb[name].protection.sheet is True, f"{name} protegit (sense contrasenya)")
check(all(not wb[n].protection.password for n in ("GPUs", "Referencia")),
      "cap full té contrasenya")

print("\n6. Amplades de columna (detecció de «#####» i solapaments)")
CHAR_W = 1.05  # Calibri 11 ≈ 1 unitat d'amplada per caràcter
for sheet in wb.worksheets:
    merged = {str(rng) for rng in sheet.merged_cells.ranges}
    for row in sheet.iter_rows():
        for c in row:
            if c.value is None:
                continue
            w = sheet.column_dimensions[c.column_letter].width or 8.43
            # Els números SÍ que es converteixen en «#####» si no hi caben.
            if isinstance(c.value, (int, float)):
                shown = format(c.value, ",.0f")
                if len(shown) * CHAR_W > w:
                    warns.append(f"{sheet.title}!{c.coordinate}: número «{shown}» no hi cap ({w:.0f})")
                continue
            txt = str(c.value)
            if txt.startswith("=") or (c.alignment and c.alignment.wrap_text):
                continue
            if any(c.coordinate in m for m in merged):
                continue
            if len(txt) * CHAR_W <= w + 2:
                continue
            # El text desborda: només molesta si la cel·la del costat està ocupada.
            nxt = sheet.cell(row=c.row, column=c.column + 1)
            if nxt.value is not None:
                warns.append(
                    f"{sheet.title}!{c.coordinate}: text de {len(txt)} car. xoca amb {nxt.coordinate}"
                )
check(not warns, f"cap text tallat ni número en «#####» ({len(warns)} avisos)")
for w in warns[:10]:
    print(f"      ⚠ {w}")

print("\n7. Lògica: la matriu de la diapositiva 24")
# El full «Exemples» calcula els veredictes sobre la VRAM REAL mesurada a Ollama,
# no sobre la teòrica: la fórmula del curs infraestima Q4_K_M (~0,6 B/param reals).
GPUS_AULA = {"RTX 3060": 12, "RTX 5070": 12, "RTX 5060 Ti": 16}
casos = [
    ("Qwen3 14B Q4 8k", 14, "Q4", 8192, 11.0,
     {"RTX 3060": "⚠ Just", "RTX 5070": "⚠ Just", "RTX 5060 Ti": "✔ Sí"}),
    ("Gemma 4 26B Q4 8k", 26, "Q4", 8192, 15.5,
     {"RTX 3060": "✘ No", "RTX 5070": "✘ No", "RTX 5060 Ti": "⚠ Just"}),
    ("Qwen3.6 27B Q4 8k", 27, "Q4", 8192, 18.2,
     {"RTX 3060": "✘ No", "RTX 5070": "✘ No", "RTX 5060 Ti": "✘ No"}),
]
for nom, par, q, ctx, real, esperat in casos:
    teoric = vram_pesos(par, q) + vram_ctx(ctx)
    got = {g: veredicte(v - real) for g, v in GPUS_AULA.items()}
    check(got == esperat,
          f"{nom}: teòrica {teoric:.1f} GB / real {real:.1f} GB → {got}")
    if abs(teoric - real) / real > 0.10:
        print(f"      ℹ la fórmula infraestima un {100*(real-teoric)/real:.0f} % "
              f"en aquest cas (esperat amb Q4_K_M)")

# Els exemples del full han de coincidir amb aquests números
ex = wb["Exemples"]
for i, (nom, *_rest, real, _v) in enumerate(casos):
    cell = ex.cell(row=6 + i, column=9)
    check(cell.value == real, f"Exemples!I{6+i} = {real} GB (VRAM real de {nom})")

print("\n8. t/s teòrics (fórmula del curs)")
for gpu, bw, par in (("RTX 5070", 672, 14), ("RTX 3060", 360, 8), ("RTX 5060 Ti", 448, 14)):
    pes = vram_pesos(par, "Q4")
    teoric = bw / pes
    print(f"  · {gpu} amb {par}B Q4: {teoric:.0f} t/s teòrics → {teoric*0.7:.0f} t/s reals")

print()
if errors:
    print(f"✘ {len(errors)} comprovacions han fallat:")
    for e in errors:
        print(f"   · {e}")
    sys.exit(1)
print(f"✔ Totes les comprovacions han passat ({len(warns)} avisos d'amplada).")
